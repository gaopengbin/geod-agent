"""Generate independent encrypted fixtures; keep passwords out of public evidence."""
from pathlib import Path
import hashlib, json, os, secrets, shutil, subprocess
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

repo=Path(__file__).resolve().parents[1]
root=repo/'artifacts/product-gaps-20261004/encrypted-documents'
workspace=root/'workspace'
workspace.mkdir(parents=True,exist_ok=True)
assert not (root/'fixtures.json').exists(),'Do not overwrite accepted encrypted fixtures'
markers={key:'CRYPTDATA'+secrets.token_hex(8).upper() for key in ['pdf','word','sheet','slides','notes','doc','xls','ppt']}
passwords={name:'QA_'+secrets.token_hex(20) for name in ['binary.doc','binary.xls','binary.ppt',*['pdf-'+a+'.pdf' for a in ['RC4-40','RC4-128','AES-128','AES-256-R5','AES-256']], 'scan-aes256.pdf','owner-only.pdf']}
for extension in ['docx','xlsx','pptx']:
    for mode in ['agile','standard']: passwords[mode+'.'+extension]='本机口令_'+secrets.token_hex(16)
passwords['pdf-AES-256.pdf']='本机口令_'+secrets.token_hex(16)
doc=Document();doc.add_paragraph('Encrypted reference: Beijing / 北京');doc.add_paragraph(markers['word']);doc.save(workspace/'plain.docx')
book=Workbook();sheet=book.active;sheet.title='Beijing';sheet.append(['Reference',markers['sheet']]);sheet.append([116.1,39.1,116.2,39.2]);sheet['E3']='=SUM(A2:D2)';book.save(workspace/'plain.xlsx')
slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='Encrypted reference';slide.placeholders[1].text=markers['slides'];slide.notes_slide.notes_text_frame.text=markers['notes'];slides.save(workspace/'plain.pptx')
pdf=canvas.Canvas(str(workspace/'plain.pdf'));pdf.drawString(60,750,'GeoD actual encrypted document acceptance');pdf.showPage();pdf.drawString(60,750,markers['pdf']);pdf.save()
runtime=repo/'apps/geod-agent-desktop/src-tauri/resources/legacy-office';jar=runtime/'tika-app-3.3.2.jar'
classes=root/'fixture-generator';classes.mkdir(exist_ok=True)
subprocess.run([shutil.which('javac'),'--release','17','-encoding','UTF-8','-cp',str(jar),'-d',str(classes),str(repo/'scripts/EncryptedOfficeFixtures.java')],check=True,creationflags=subprocess.CREATE_NO_WINDOW)
subprocess.run([str(runtime/'java/bin/java.exe'),'-Dfile.encoding=UTF-8','-cp',str(classes)+os.pathsep+str(jar),'EncryptedOfficeFixtures',str(workspace),str(repo/'artifacts/product-gaps-20261004/legacy-office/apache-empty.doc')],input=json.dumps(dict(passwords=passwords,markers=markers),ensure_ascii=False).encode('utf-8'),capture_output=True,check=True,creationflags=subprocess.CREATE_NO_WINDOW)
code='''import sys,json,hashlib
from pathlib import Path
sys.path.insert(0,sys.argv[2])
import cryptography,cffi,_cffi_backend,olefile,msoffcrypto
from cryptography.hazmat.primitives.ciphers import Cipher,algorithms,modes
from pypdf import PdfReader,PdfWriter
p=json.load(sys.stdin);root=Path(sys.argv[1])
for algorithm in ['RC4-40','RC4-128','AES-128','AES-256-R5','AES-256']:
  name='pdf-'+algorithm+'.pdf';writer=PdfWriter();writer.append(PdfReader(root/'plain.pdf'));writer.encrypt(p[name],owner_password=p[name]+'-owner',algorithm=algorithm);writer.write(root/name)
scan=PdfWriter();scan.append(PdfReader(sys.argv[3]));scan.encrypt(p['scan-aes256.pdf'],algorithm='AES-256');scan.write(root/'scan-aes256.pdf')
owner=PdfWriter();owner.append(PdfReader(root/'plain.pdf'));owner.encrypt('',owner_password=p['owner-only.pdf'],algorithm='AES-256');owner.write(root/'owner-only.pdf')
sample=root/'agile.docx'
with sample.open('rb') as f:
  data=msoffcrypto.OfficeFile(f);data.load_key(password=p['agile.docx'],verify_password=True)
  assert data.is_encrypted()
print(json.dumps(dict(python=sys.version.split()[0],cryptography=cryptography.__version__,cffi=cffi.__version__,nativeCffi=_cffi_backend.__version__,office=msoffcrypto.__file__)))
'''
env={key:value for key,value in os.environ.items() if key.upper() in ['SYSTEMROOT','WINDIR','TEMP','TMP']}
python=repo/'apps/geod-agent-desktop/src-tauri/resources/gdal/python.exe'
completed=subprocess.run([str(python),'-I','-X','utf8','-c',code,str(workspace),str(repo/'apps/geod-agent-desktop/src-tauri/resources/documents'),str(repo/'artifacts/product-gaps-20261004/document-ocr/workspace/scan-en.pdf')],input=json.dumps(passwords,ensure_ascii=False).encode('utf-8'),env=env,capture_output=True,timeout=60,creationflags=subprocess.CREATE_NO_WINDOW)
assert completed.returncode==0,completed.stderr.decode('utf-8',errors='replace')
(root/'runtime-imports.json').write_bytes(completed.stdout)
files={name:{'path':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'bytes':file.stat().st_size} for name in passwords if (file:=workspace/name).is_file()}
assert len(files)==len(passwords)==16
for name in [*['agile.'+ext for ext in ['docx','xlsx','pptx']],*['standard.'+ext for ext in ['docx','xlsx','pptx']],*['binary.'+ext for ext in ['doc','xls','ppt']]]: assert (workspace/name).read_bytes()[:8]==bytes.fromhex('d0cf11e0a1b11ae1')
scan_marker=json.loads((repo/'artifacts/product-gaps-20261004/document-ocr/fixtures.json').read_text(encoding='utf-8'))['markers']['en']
(root/'fixtures.json').write_text(json.dumps(dict(authoredOnly=True,markers=markers,scanMarker=scan_marker,files=files,independentOfficeWriter='Apache POI 5.5.1'),ensure_ascii=False,indent=2),encoding='utf-8')
# QA-only secret is removed before the persistence audit, never stored in the app.
(root/'private-passwords.json').write_text(json.dumps(passwords,ensure_ascii=False),encoding='utf-8')
print(json.dumps(dict(prepared=True,encryptedFiles=len(files),passwordsInPublicEvidence=False),ensure_ascii=False))
