"""Fresh encrypted inputs and verified authored samples for the integrated release."""
from pathlib import Path
import hashlib
import json
import os
import secrets
import shutil
import subprocess
from docx import Document
from reportlab.pdfgen import canvas

repo = Path(__file__).resolve().parents[1]
root = repo/'artifacts/product-gaps-20261004/integrated-release'
workspace = root/'workspace'
assert (root/'postgis-ready.json').is_file()
assert not (root/'fixtures.json').exists(), 'Keep the authored fixtures unchanged'
workspace.mkdir(exist_ok=True)
passwords = {name: '本机🗝_'+secrets.token_hex(24) for name in ['release-encrypted.pdf', 'release-encrypted.docx']}
markers = {name: 'PACKDATA'+secrets.token_hex(10).upper() for name in ['pdf', 'word', 'audio', 'sqlite']}
doc = Document()
doc.add_paragraph('Integrated GeoD release document acceptance')
doc.add_paragraph(markers['word'])
doc.save(workspace/'release-plain.docx')
pdf = canvas.Canvas(str(workspace/'release-plain.pdf'))
pdf.drawString(60, 750, 'Integrated GeoD release PDF acceptance')
pdf.drawString(60, 720, markers['pdf'])
pdf.save()
code = '''
import sys,json
from pathlib import Path
sys.path.insert(0,sys.argv[2])
from pypdf import PdfReader,PdfWriter
import msoffcrypto
root=Path(sys.argv[1]);passwords=json.load(sys.stdin)
writer=PdfWriter();writer.append(PdfReader(root/'release-plain.pdf'))
writer.encrypt(passwords['release-encrypted.pdf'],algorithm='AES-256')
writer.write(root/'release-encrypted.pdf')
with (root/'release-plain.docx').open('rb') as source,(root/'release-encrypted.docx').open('wb') as target:
    msoffcrypto.OfficeFile(source).encrypt(passwords['release-encrypted.docx'],target)
'''
env = {k: v for k, v in os.environ.items() if k.upper() in ['SYSTEMROOT', 'WINDIR', 'TEMP', 'TMP']}
python = repo/'apps/geod-agent-desktop/src-tauri/resources/gdal/python.exe'
completed = subprocess.run([str(python), '-I', '-X', 'utf8', '-c', code, str(workspace), str(repo/'apps/geod-agent-desktop/src-tauri/resources/documents')], input=json.dumps(passwords, ensure_ascii=False).encode(), env=env, capture_output=True, timeout=60, creationflags=subprocess.CREATE_NO_WINDOW)
assert completed.returncode == 0, 'The pinned fixture generator failed; no password diagnostics are printed'
(root/'private/document-passwords.json').write_text(json.dumps(passwords, ensure_ascii=False), encoding='utf-8')
legacy = json.loads((repo/'artifacts/product-gaps-20261004/legacy-office/fixtures.json').read_text(encoding='utf-8'))
for name, kind in [('beijing-brief.doc', 'en'), ('beijing-table.xls', 'xls'), ('beijing-slides.ppt', 'ppt')]:
    source = Path(legacy['files'][name]['path'])
    assert hashlib.sha256(source.read_bytes()).hexdigest() == legacy['files'][name]['sha256']
    shutil.copyfile(source, workspace/name)
    markers[kind] = legacy['markers'][kind]
ocr = json.loads((repo/'artifacts/product-gaps-20261004/document-ocr/fixtures.json').read_text(encoding='utf-8'))
source = repo/'artifacts/product-gaps-20261004/document-ocr/workspace/scan-en.pdf'
assert hashlib.sha256(source.read_bytes()).hexdigest() == ocr['files']['scan-en.pdf']['sha256']
shutil.copyfile(source, workspace/'release-scan.pdf')
markers['scan'] = ocr['markers']['en']
shutil.copyfile(repo/'artifacts/product-gaps-20261004/audio-inputs/workspace/beijing-english.mp3', workspace/'release-Beijing.mp3')
import sqlite3
with sqlite3.connect(workspace/'release.sqlite') as db:
    db.execute('CREATE TABLE release_markers(marker TEXT,city TEXT)')
    db.execute('INSERT INTO release_markers VALUES(?,?)', (markers['sqlite'], 'Beijing'))
inputs = ['release-encrypted.pdf', 'release-encrypted.docx', 'release-scan.pdf', 'beijing-brief.doc', 'beijing-table.xls', 'beijing-slides.ppt', 'release-Beijing.mp3']
files = {name: dict(sha256=hashlib.sha256((workspace/name).read_bytes()).hexdigest(), bytes=(workspace/name).stat().st_size) for name in inputs}
(root/'fixtures.json').write_text(json.dumps(dict(markers=markers, inputs=inputs, files=files, audioMarkerIsUserCorrection=True), ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(dict(prepared=True, inputs=len(inputs), encryptedInputs=len(passwords), sqlite=True)), flush=True)
