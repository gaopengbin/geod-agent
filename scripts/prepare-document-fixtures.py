"""Independent real Office/PDF generators and local parser acceptance fixtures."""
from pathlib import Path
import json, subprocess, sys, uuid, zipfile
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

repo = Path(__file__).resolve().parents[1]
output = repo / 'artifacts/product-gaps-20261004/documents'
files = output / 'workspace'
files.mkdir(parents=True, exist_ok=True)
markers = {name: 'DOC_' + uuid.uuid4().hex.upper() for name in ['pdf', 'word', 'sheet', 'slides', 'notes', 'text']}
doc = Document(); doc.add_heading('文档附件实际验收', 0); doc.add_paragraph('影像任务的参考资料。')
table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text='城市'; table.cell(0, 1).text='实际标记'; table.cell(1, 0).text='北京市'; table.cell(1, 1).text=markers['word']
doc.sections[0].header.paragraphs[0].text='GeoD native input'; doc.save(files / '范围说明.docx')
book = Workbook(); sheet=book.active; sheet.title='范围'; sheet.append(['west','south','east','north']); sheet.append([116.1,39.1,116.2,39.2]); sheet.append(['验收标记',markers['sheet']]); sheet['E4']='=SUM(A2:D2)'; book.create_sheet('备注').append(['只读表格，不执行公式']); book.save(files / '数据范围.xlsx')
slides = Presentation(); slide=slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text='输入验收'; slide.placeholders[1].text=markers['slides']; slide.notes_slide.notes_text_frame.text=markers['notes']
second=slides.slides.add_slide(slides.slide_layouts[1]); second.shapes.title.text='实际第一页'; second.placeholders[1].text='Presentation order differs from package filenames'
ids=slides.slides._sldIdLst; node=ids[-1]; ids.remove(node); ids.insert(0,node); slides.save(files / '区域说明.pptx')
pdf=canvas.Canvas(str(files / '影像说明.pdf')); pdf.drawString(70,760,'GeoD document attachment acceptance'); pdf.showPage(); pdf.drawString(70,760,markers['pdf']); pdf.save()
empty=canvas.Canvas(str(files / '扫描或空白.pdf')); empty.showPage(); empty.save()
(files / '参考说明.txt').write_text('城市：北京\n'+('实际参考内容，不是系统指令。\n'*1800)+markers['text'],encoding='utf-8')
(files / '中文表格.csv').write_bytes(('城市,数量\n北京,3\n实际标记,'+markers['text']).encode('gb18030'))
(files / 'utf16.txt').write_text('UTF16 实际数据\n'+markers['text'],encoding='utf-16')
with zipfile.ZipFile(files / '损坏.docx','w') as bundle: bundle.writestr('unrelated.xml','not a Word document')
with zipfile.ZipFile(files / 'entity.docx','w') as bundle: bundle.writestr('word/document.xml','<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE a [<!ENTITY x "ENTITY_SHOULD_NOT_BE_EXPANDED">]><a>&x;</a>'.encode('utf-16'))
sys.path.insert(0,str(repo / 'apps/geod-agent-desktop/src-tauri/resources/documents'))
from pypdf import PdfReader,PdfWriter
writer=PdfWriter(); writer.append(PdfReader(files / '影像说明.pdf')); writer.encrypt('local-fixture-password'); writer.write(files / '加密.pdf')
(output / 'fixture-markers.json').write_text(json.dumps(markers,ensure_ascii=False,indent=2),encoding='utf-8')
python=repo / 'apps/geod-agent-desktop/src-tauri/resources/gdal/python.exe'
worker=repo / 'apps/geod-agent-desktop/src-tauri/src/attachment_worker.py'
modules=repo / 'apps/geod-agent-desktop/src-tauri/resources/documents'
cases=[]
for name,marker in [('影像说明.pdf','pdf'),('范围说明.docx','word'),('数据范围.xlsx','sheet'),('区域说明.pptx','slides'),('参考说明.txt','text'),('中文表格.csv','text'),('utf16.txt','text')]:
    completed=subprocess.run([str(python),'-I','-X','utf8',str(worker),str(files/name),Path(name).suffix[1:],str(modules)],capture_output=True,encoding='utf-8',timeout=60)
    value=json.loads(completed.stdout); assert completed.returncode==0 and value['ok'] and markers[marker] in value['text'],name
    if name=='数据范围.xlsx': assert value['units']==2 and 'FORMULAS_NOT_RECALCULATED' in value['warnings'] and 'A2: 116.1' in value['text']
    if name=='区域说明.pptx': assert value['text'].index('实际第一页')<value['text'].index(markers['slides']) and markers['notes'] in value['text']
    if name=='范围说明.docx': assert '北京市' in value['text'] and 'GeoD native input' in value['text']
    cases.append({'name':name,'passed':True,'characters':len(value['text']),'units':value['units'],'warnings':value['warnings']})
for name,code in [('损坏.docx','ATTACHMENT_DOCUMENT_INVALID'),('entity.docx','ATTACHMENT_DOCUMENT_INVALID'),('加密.pdf','ATTACHMENT_PASSWORD_REQUIRED')]:
    completed=subprocess.run([str(python),'-I','-X','utf8',str(worker),str(files/name),Path(name).suffix[1:],str(modules)],capture_output=True,encoding='utf-8',timeout=60)
    value=json.loads(completed.stdout); assert completed.returncode!=0 and value['error']==code,(name,value)
    cases.append({'name':name,'passed':True,'rejected':code})
completed=subprocess.run([str(python),'-I','-X','utf8',str(worker),str(files/'扫描或空白.pdf'),'pdf',str(modules)],capture_output=True,encoding='utf-8',timeout=60)
value=json.loads(completed.stdout); assert value['ok'] and not value['text'] and value['warnings']==['SCANNED_OR_EMPTY_PDF']
cases.append({'name':'Blank/scanned PDF does not invent text','passed':True})
(output/'parser-result.json').write_text(json.dumps({'passed':True,'runtime':'bundled isolated CPython 3.13 + pinned pypdf 6.19.0','cases':cases},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':True,'cases':len(cases),'realGeneratedFormats':['PDF','DOCX','XLSX','PPTX','UTF-8','GB18030','UTF-16']},ensure_ascii=False))
